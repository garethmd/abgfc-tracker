"""Spreadsheet export: the season as the coaches' original Google Sheet laid it out.

Five tabs - Summary, Fixtures, Match Stats, Appearances, Squad - with the same headers,
colours, validations and formulas, so a coach who liked the sheet gets the sheet. The
data tabs hold what the app knows; the Summary recalculates from them in Excel, so the
file keeps working if someone edits it offline (it will not sync back).

Where the original used Google-only functions (ARRAYFORMULA / FILTER / TEXTJOIN for the
scorers column, last-5 form and highlight names) the app writes the computed text instead
and says so in a cell comment.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from app.core.errors import NotFoundError
from app.models import CompetitionType, EventType, Fixture, FixtureStatus, UserRole
from app.repositories.club import TeamSeasonRepository
from app.repositories.fixtures import FixtureRepository
from app.repositories.players import SquadRepository
from app.services.access import Access
from app.services.stats import StatsService, form

NAVY = "1A3670"
CALC = "EEF0F3"  # calculated cells
TOTAL = "E0E9F5"  # summary totals
GREY = "666666"
FONT = "Arial"

FIXTURE_ROWS = 40  # Fixtures!2:41, as in the original
STATS_ROWS = 500  # Match Stats!2:501
SQUAD_ROWS = 30  # Squad!2:31 -> up to 27 player columns on Appearances (D..AD)
SUMMARY_SQUAD_ROWS = 20  # Summary!16:35

COMP_LABEL = {
    CompetitionType.LEAGUE: "League",
    CompetitionType.CUP: "Cup",
    CompetitionType.FRIENDLY: "Friendly",
    CompetitionType.TOURNAMENT: "Tournament",
}


# --- styling helpers -----------------------------------------------------------------


def _header(ws, row: int, labels: list[str]) -> None:
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)


def _calc(cell, fill: str = CALC, bold: bool = False, fmt: str | None = None):
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name=FONT, bold=bold)
    if fmt:
        cell.number_format = fmt
    return cell


def _title(ws, coord: str, text: str, size: int | None = None):
    c = ws[coord]
    c.value = text
    c.font = Font(name=FONT, bold=True, color=NAVY, size=size)
    return c


def _note(ws, coord: str, text: str):
    c = ws[coord]
    c.value = text
    c.font = Font(name=FONT, color=GREY)
    return c


def _widths(ws, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _plain(ws):
    """Arial everywhere a font hasn't been set explicitly."""
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font.name != FONT:
                c.font = Font(name=FONT, bold=c.font.bold, color=c.font.color, size=c.font.size)


# --- the workbook --------------------------------------------------------------------


def build_workbook(db: Session, access: Access, team_season_id: int) -> tuple[bytes, str]:
    ts = TeamSeasonRepository(db).get(team_season_id)
    if ts is None:
        raise NotFoundError(f"Team season {team_season_id} not found")
    access.require_team_season(ts, UserRole.COACH)  # names children: coaches only

    team, season = ts.club_team, ts.season
    fixtures = FixtureRepository(db).list_all(team_season_id=ts.id)
    members = SquadRepository(db).list_for_team_season(ts.id)
    # Squad order: number, then name - the same order the app shows.
    players = [m.player for m in members]
    details = {m.player_id: m for m in members}
    played = [f for f in fixtures if f.status == FixtureStatus.PLAYED]
    rows = StatsService(db, access).leaderboard(ts.id).rows
    fixture_detail = {f.id: f for f in fixtures}

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_fix = wb.create_sheet("Fixtures")
    ws_stats = wb.create_sheet("Match Stats")
    ws_app = wb.create_sheet("Appearances")
    ws_squad = wb.create_sheet("Squad")

    _squad_sheet(ws_squad, players, details)
    _fixtures_sheet(ws_fix, fixtures, fixture_detail)
    _match_stats_sheet(ws_stats, fixtures)
    _appearances_sheet(ws_app, fixtures, players)
    _summary_sheet(ws_summary, team.name, season.name, played, rows)

    for ws in wb.worksheets:
        _plain(ws)

    buf = BytesIO()
    wb.save(buf)
    stamp = datetime.now().strftime("%Y-%m-%d")
    return buf.getvalue(), f"ABGFC-{team.name}-{season.name.replace('/', '-')}-{stamp}.xlsx"


def _squad_sheet(ws, players, details) -> None:
    _header(ws, 1, ["Player", "Notes (optional)"])
    _note(ws, "C1", "Exported from the ABGFC app - add or remove players there.")
    for i, p in enumerate(players, start=2):
        ws.cell(row=i, column=1, value=p.display_name)
        m = details[p.id]
        bits = []
        if m.squad_number is not None:
            bits.append(f"#{m.squad_number}")
        if m.primary_position:
            bits.append(m.primary_position.name)
        if m.left_at:
            bits.append(f"left {m.left_at.isoformat()}")
        if bits:
            ws.cell(row=i, column=2, value=" · ".join(bits))
    ws.freeze_panes = "A2"
    _widths(ws, {"A": 18.9, "B": 25.1, "C": 50.1})


def _fixtures_sheet(ws, fixtures: list[Fixture], detail) -> None:
    _header(
        ws,
        1,
        [
            "Match #",
            "Date",
            "Competition",
            "Opposition",
            "Our Score",
            "Their Score",
            "Result",
            "Scorers",
            "Player of the Match",
            "Parents' Player of the Match",
            "Notes (venue, kick-off, etc.)",
        ],
    )
    by_number = {f.match_number: f for f in fixtures if f.match_number}
    unnumbered = [f for f in fixtures if not f.match_number]
    for r in range(2, FIXTURE_ROWS + 2):
        n = r - 1
        ws.cell(row=r, column=1, value=n)
        f = by_number.get(n) or (unnumbered.pop(0) if unnumbered else None)
        if f is not None:
            d = ws.cell(row=r, column=2, value=f.kickoff_at.date())
            d.number_format = "ddd dd mmm yyyy"
            ws.cell(row=r, column=3, value=COMP_LABEL[CompetitionType(f.competition.type)])
            ws.cell(row=r, column=4, value=f.opposition.name)
            if f.status == FixtureStatus.PLAYED:
                ws.cell(row=r, column=5, value=f.our_score)
                ws.cell(row=r, column=6, value=f.their_score)
            notes = [
                f.competition.name,
                {"home": "Home", "away": "Away", "neutral": "Neutral"}[f.venue],
            ]
            if f.venue_notes:
                notes.append(f.venue_notes)
            notes.append(f.kickoff_at.strftime("%H:%M"))
            if f.status not in (FixtureStatus.PLAYED, FixtureStatus.SCHEDULED):
                notes.append(f.status.upper())
            if f.notes:
                notes.append(f.notes)
            ws.cell(row=r, column=11, value=" · ".join(notes))
        _calc(
            ws.cell(
                row=r,
                column=7,
                value=f'=IF(OR($E{r}="",$F{r}=""),"",IF($E{r}>$F{r},"W",IF($E{r}=$F{r},"D","L")))',
            ),
            bold=True,
        )
        scorers, coaches, parents = _fixture_text(f) if f else ("", "", "")
        _calc(ws.cell(row=r, column=8, value=scorers))
        _calc(ws.cell(row=r, column=9, value=coaches))
        _calc(ws.cell(row=r, column=10, value=parents))
    ws["H1"].comment = Comment(
        "Scorers, POTM and Parents' POTM are written by the app at export time "
        "(the original sheet used Google-only FILTER/TEXTJOIN formulas).",
        "ABGFC",
    )
    ws.conditional_formatting.add(
        f"G2:G{FIXTURE_ROWS + 1}",
        CellIsRule(
            operator="equal",
            formula=['"W"'],
            fill=PatternFill("solid", fgColor="D3EDD6"),
            font=Font(color="1C5E21", bold=True),
        ),
    )
    ws.conditional_formatting.add(
        f"G2:G{FIXTURE_ROWS + 1}",
        CellIsRule(
            operator="equal",
            formula=['"D"'],
            fill=PatternFill("solid", fgColor="FFF2CC"),
            font=Font(color="7A5900", bold=True),
        ),
    )
    ws.conditional_formatting.add(
        f"G2:G{FIXTURE_ROWS + 1}",
        CellIsRule(
            operator="equal",
            formula=['"L"'],
            fill=PatternFill("solid", fgColor="F9D8D8"),
            font=Font(color="A51C1C", bold=True),
        ),
    )
    dv = DataValidation(type="list", formula1='"League,Cup,Friendly,Tournament"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C2:C{FIXTURE_ROWS + 1}")
    ws.freeze_panes = "A2"
    _widths(
        ws,
        {
            "A": 8.9,
            "B": 16.4,
            "C": 13.2,
            "D": 21.4,
            "E": 12.0,
            "F": 12.0,
            "G": 8.9,
            "H": 30.1,
            "I": 18.9,
            "J": 22.0,
            "K": 40.0,
        },
    )


def _fixture_text(f: Fixture) -> tuple[str, str, str]:
    if f.status != FixtureStatus.PLAYED:
        return "", "", ""
    goals: dict[str, int] = {}
    for e in f.events:
        if e.event_type == EventType.GOAL and e.player:
            goals[e.player.display_name] = goals.get(e.player.display_name, 0) + 1
    scorers = ", ".join(f"{n} ({c})" if c > 1 else n for n, c in goals.items())
    coaches = ", ".join(
        a.player.display_name for a in f.awards if a.award_type.code == "coaches_potm"
    )
    parents = ", ".join(
        a.player.display_name for a in f.awards if a.award_type.code == "parents_potm"
    )
    return scorers, coaches, parents


def _match_stats_sheet(ws, fixtures: list[Fixture]) -> None:
    _header(ws, 1, ["Match #", "Opposition", "Player", "Goals", "Assists", "POTM", "Parents' POTM"])
    r = 2
    for f in fixtures:
        if f.status != FixtureStatus.PLAYED:
            continue
        per = _per_player(f)
        for d in per.values():
            ws.cell(row=r, column=1, value=f.match_number)
            ws.cell(row=r, column=3, value=d["name"])
            if d["g"]:
                ws.cell(row=r, column=4, value=d["g"])
            if d["a"]:
                ws.cell(row=r, column=5, value=d["a"])
            if d["c"]:
                ws.cell(row=r, column=6, value="Y")
            if d["p"]:
                ws.cell(row=r, column=7, value="Y")
            r += 1
    for rr in range(2, STATS_ROWS + 2):
        _calc(
            ws.cell(
                row=rr,
                column=2,
                value=f'=IF($A{rr}="","",IFERROR(VLOOKUP($A{rr},Fixtures!$A$2:$D${FIXTURE_ROWS + 1},4,FALSE),""))',
            )
        )
    yes = DataValidation(type="list", formula1='"Y"', allow_blank=True)
    ws.add_data_validation(yes)
    yes.add(f"F2:G{STATS_ROWS + 1}")
    who = DataValidation(type="list", formula1=f"Squad!$A$2:$A${SQUAD_ROWS + 1}", allow_blank=True)
    ws.add_data_validation(who)
    who.add(f"C2:C{STATS_ROWS + 1}")
    which = DataValidation(
        type="list", formula1=f"Fixtures!$A$2:$A${FIXTURE_ROWS + 1}", allow_blank=True
    )
    ws.add_data_validation(which)
    which.add(f"A2:A{STATS_ROWS + 1}")
    ws.freeze_panes = "A2"
    _widths(ws, {"A": 8.9, "B": 20.1, "C": 16.4, "D": 8.9, "E": 8.9, "F": 8.9, "G": 13.9})


def _per_player(f: Fixture) -> dict[int, dict]:
    """One Match Stats row per player who scored, assisted or won an award."""
    per: dict[int, dict] = {}

    def row(player):
        return per.setdefault(
            player.id, {"name": player.display_name, "g": 0, "a": 0, "c": "", "p": ""}
        )

    for e in f.events:
        if e.player is None:
            continue
        if e.event_type == EventType.GOAL:
            row(e.player)["g"] += 1
        elif e.event_type == EventType.ASSIST:
            row(e.player)["a"] += 1
    for a in f.awards:
        if a.award_type.code == "coaches_potm":
            row(a.player)["c"] = "Y"
        elif a.award_type.code == "parents_potm":
            row(a.player)["p"] = "Y"
    return per


def _appearances_sheet(ws, fixtures: list[Fixture], players) -> None:
    _header(ws, 1, ["Match #", "Opposition", "# Played"])
    last_col = 3 + SQUAD_ROWS - 3  # D..AD = 27 player columns, as the original
    for i in range(last_col - 3):
        c = ws.cell(row=1, column=4 + i, value=f'=IF(Squad!$A{2 + i}="","",Squad!$A{2 + i})')
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
    col_of = {p.id: 4 + i for i, p in enumerate(players)}
    by_number = {f.match_number: f for f in fixtures if f.match_number}
    end = get_column_letter(last_col)
    for r in range(2, FIXTURE_ROWS + 2):
        n = r - 1
        ws.cell(row=r, column=1, value=n)
        _calc(
            ws.cell(
                row=r,
                column=2,
                value=f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},Fixtures!$A$2:$D${FIXTURE_ROWS + 1},4,FALSE),""))',
            )
        )
        _calc(ws.cell(row=r, column=3, value=f'=IF($B{r}="","",COUNTIF($D{r}:${end}{r},"Y"))'))
        f = by_number.get(n)
        if f is not None and f.status == FixtureStatus.PLAYED:
            for a in f.appearances:
                col = col_of.get(a.player_id)
                if col:
                    ws.cell(row=r, column=col, value="Y")
    yes = DataValidation(type="list", formula1='"Y"', allow_blank=True)
    ws.add_data_validation(yes)
    yes.add(f"D2:{end}{FIXTURE_ROWS + 1}")
    ws.freeze_panes = "D2"
    _widths(ws, {"A": 8.9, "B": 20.1, "C": 9.5})
    for i in range(4, last_col + 1):
        ws.column_dimensions[get_column_letter(i)].width = 10.8


def _summary_sheet(ws, team_name: str, season_name: str, played: list[Fixture], rows) -> None:
    FR = FIXTURE_ROWS + 1
    SR = STATS_ROWS + 1
    _title(ws, "A1", f"ABGFC {team_name} — {season_name} Season Summary", size=16)
    _note(
        ws,
        "A2",
        "Exported from the ABGFC app. The grey cells recalculate from the other tabs; edits here don't sync back to the app.",
    )

    _title(ws, "A4", "SEASON RECORD — ALL MATCHES")
    _header(
        ws,
        5,
        ["Played", "Won", "Drawn", "Lost", "Goals For", "Goals Against", "Goal Diff", "Win %"],
    )
    _calc(ws["A6"], TOTAL, True).value = "=$B6+$C6+$D6"
    _calc(ws["B6"], TOTAL, True).value = f'=COUNTIF(Fixtures!$G$2:$G${FR},"W")'
    _calc(ws["C6"], TOTAL, True).value = f'=COUNTIF(Fixtures!$G$2:$G${FR},"D")'
    _calc(ws["D6"], TOTAL, True).value = f'=COUNTIF(Fixtures!$G$2:$G${FR},"L")'
    _calc(ws["E6"], TOTAL, True).value = f"=SUM(Fixtures!$E$2:$E${FR})"
    _calc(ws["F6"], TOTAL, True).value = f"=SUM(Fixtures!$F$2:$F${FR})"
    _calc(ws["G6"], TOTAL, True).value = "=$E6-$F6"
    _calc(ws["H6"], TOTAL, True, "0%").value = '=IF($A6=0,"",$B6/$A6)'

    _title(ws, "A8", "LEAGUE MATCHES ONLY")
    _header(
        ws,
        9,
        ["Played", "Won", "Drawn", "Lost", "Goals For", "Goals Against", "Goal Diff", "Win %"],
    )
    _calc(ws["A10"], TOTAL, True).value = "=$B10+$C10+$D10"
    for col, res in (("B", "W"), ("C", "D"), ("D", "L")):
        _calc(
            ws[f"{col}10"], TOTAL, True
        ).value = f'=COUNTIFS(Fixtures!$G$2:$G${FR},"{res}",Fixtures!$C$2:$C${FR},"League")'
    _calc(
        ws["E10"], TOTAL, True
    ).value = f'=SUMIFS(Fixtures!$E$2:$E${FR},Fixtures!$C$2:$C${FR},"League")'
    _calc(
        ws["F10"], TOTAL, True
    ).value = f'=SUMIFS(Fixtures!$F$2:$F${FR},Fixtures!$C$2:$C${FR},"League")'
    _calc(ws["G10"], TOTAL, True).value = "=$E10-$F10"
    _calc(ws["H10"], TOTAL, True, "0%").value = '=IF($A10=0,"",$B10/$A10)'

    _title(ws, "A12", "Last 5 results (most recent first)")
    last5 = [x.result for x in reversed(form(played))]
    c = _calc(ws["C12"], bold=True)
    c.value = "   ".join(last5) if last5 else "—"
    c.comment = Comment(
        "Written by the app at export time (Google-only formula in the original).", "ABGFC"
    )

    _title(ws, "A14", "SQUAD")
    _header(
        ws, 15, ["Player", "Pld", "Goals", "Assists", "Goals per game", "POTM", "Parents' POTM"]
    )
    for i in range(SUMMARY_SQUAD_ROWS):
        r = 16 + i
        _calc(ws[f"A{r}"]).value = f'=IF(Squad!$A{2 + i}="","",Squad!$A{2 + i})'
        _calc(ws[f"B{r}"]).value = (
            f'=IF($A{r}="","",IFERROR(COUNTIF(INDEX(Appearances!$D$2:$AD${FR},0,'
            f'MATCH($A{r},Appearances!$D$1:$AD$1,0)),"Y"),0))'
        )
        _calc(
            ws[f"C{r}"]
        ).value = f"=IF($A{r}=\"\",\"\",SUMIF('Match Stats'!$C$2:$C${SR},$A{r},'Match Stats'!$D$2:$D${SR}))"
        _calc(
            ws[f"D{r}"]
        ).value = f"=IF($A{r}=\"\",\"\",SUMIF('Match Stats'!$C$2:$C${SR},$A{r},'Match Stats'!$E$2:$E${SR}))"
        _calc(
            ws[f"E{r}"], fmt="0.00"
        ).value = f'=IF($A{r}="","",IF($B{r}=0,"",ROUND($C{r}/$B{r},2)))'
        _calc(
            ws[f"F{r}"]
        ).value = f'=IF($A{r}="","",COUNTIFS(\'Match Stats\'!$C$2:$C${SR},$A{r},\'Match Stats\'!$F$2:$F${SR},"Y"))'
        _calc(
            ws[f"G{r}"]
        ).value = f'=IF($A{r}="","",COUNTIFS(\'Match Stats\'!$C$2:$C${SR},$A{r},\'Match Stats\'!$G$2:$G${SR},"Y"))'

    _title(ws, "A37", "HIGHLIGHTS")
    top = 16 + SUMMARY_SQUAD_ROWS - 1
    tiles = [
        ("Top scorer", "C", "goal", lambda r: r.goals),
        ("Most assists", "D", "assist", lambda r: r.assists),
        (
            "Most POTM (coaches)",
            "F",
            "award",
            lambda r: next((a.count for a in r.awards if a.award_type_code == "coaches_potm"), 0),
        ),
        (
            "Most POTM (parents)",
            "G",
            "award",
            lambda r: next((a.count for a in r.awards if a.award_type_code == "parents_potm"), 0),
        ),
    ]
    for i, (label, col, noun, value_of) in enumerate(tiles):
        r = 38 + i
        ws[f"A{r}"].value = label
        ws[f"A{r}"].font = Font(name=FONT, bold=True)
        best = max((value_of(x) for x in rows), default=0)
        winners = sorted(x.player.display_name for x in rows if best > 0 and value_of(x) == best)
        _calc(ws[f"C{r}"], bold=True).value = ", ".join(winners) if winners else "—"
        _calc(ws[f"D{r}"]).value = (
            f'=IF(MAX(${col}$16:${col}${top})=0,"",MAX(${col}$16:${col}${top})&" {noun}"'
            f'&IF(MAX(${col}$16:${col}${top})=1,"","s"))'
        )
    ws["C38"].comment = Comment(
        "Names written by the app at export time; the count beside them is a live formula.", "ABGFC"
    )

    _title(ws, "A43", "ABOUT THIS FILE")
    _note(
        ws,
        "A44",
        "This is a snapshot from the app. Enter results in the app, then export again for an up-to-date copy.",
    )
    _note(
        ws,
        "A45",
        "Fixtures, Match Stats, Appearances and Squad hold the data; Summary recalculates from them if you edit offline.",
    )
    _note(ws, "A46", "Grey cells are calculated — typing in them will break the totals.")
    _widths(
        ws, {"A": 23.2, "B": 10.1, "C": 18.9, "D": 16.4, "E": 13.9, "F": 10.1, "G": 13.9, "H": 10.1}
    )
    ws.column_dimensions["A"].width = 23.2
