"""The season spreadsheet: same tabs and columns as the coaches' original sheet, with
formulas that recalculate to the numbers the app reports."""

import os
import shutil
import subprocess
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.models import RoleScope, UserRole
from app.services.bootstrap import DemoSeason
from tests.conftest import login_as, make_user

API = "/api/v1"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _download(client, team_season_id):
    r = client.get(f"{API}/team-seasons/{team_season_id}/reports/season.xlsx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX
    return r.content


def test_workbook_structure_matches_the_sheet(auth_client: TestClient, demo: DemoSeason):
    data = _download(auth_client, demo.team_season.id)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Summary", "Fixtures", "Match Stats", "Appearances", "Squad"]

    fx = wb["Fixtures"]
    assert [c.value for c in fx[1]][:11] == [
        "Match #",
        "Date",
        "Competition",
        "Opposition",
        "Our Score",
        "Their Score",
        "Result",
        "Scorers",
        "Player of the Match",
        "Parents' Player of the Match",
        "Notes (venue, kick-off, etc.)",
    ]
    # Match 4: 5-0 v Hook, Archie hat-trick
    row4 = [c.value for c in fx[5]]
    assert row4[0] == 4 and row4[3] == "Hook Juniors" and (row4[4], row4[5]) == (5, 0)
    assert row4[2] == "League" and row4[6].startswith("=IF(")  # result is a live formula
    assert (
        row4[7] == "Archie (3), Ayla, Kayson" and row4[8] == "Archie" and row4[9] == "Ayla, Kayson"
    )
    # Scheduled fixture: no scores, so the result formula yields blank
    row7 = [c.value for c in fx[8]]
    assert row7[3] == "Fleet Spurs" and row7[4] is None and not row7[7]
    assert fx.freeze_panes == "A2" and fx.conditional_formatting

    ms = wb["Match Stats"]
    assert [c.value for c in ms[1]][:7] == [
        "Match #",
        "Opposition",
        "Player",
        "Goals",
        "Assists",
        "POTM",
        "Parents' POTM",
    ]
    rows = [[c.value for c in r] for r in ms.iter_rows(min_row=2, max_row=40) if r[0].value]
    archie_m4 = next(r for r in rows if r[0] == 4 and r[2] == "Archie")
    assert (archie_m4[3], archie_m4[4], archie_m4[5]) == (3, 1, "Y")
    assert all(r[1].startswith("=IF(") for r in rows)  # opposition is looked up

    ap = wb["Appearances"]
    assert [ap.cell(1, c).value for c in range(1, 4)] == ["Match #", "Opposition", "# Played"]
    assert ap["D1"].value == '=IF(Squad!$A2="","",Squad!$A2)'
    ys = [ap.cell(6, c).value for c in range(4, 31)].count("Y")  # match 5: whole squad
    assert ys == 11 and ap.freeze_panes == "D2"

    sq = wb["Squad"]
    names = [sq.cell(r, 1).value for r in range(2, 13)]
    assert names[0] == "Alexander" and len([n for n in names if n]) == 11
    assert sq["B2"].value == "#1"  # squad number carried into Notes

    sm = wb["Summary"]
    assert sm["A1"].value == "ABGFC Blues — 2026/27 Season Summary"
    assert sm["B6"].value == '=COUNTIF(Fixtures!$G$2:$G$41,"W")'
    assert sm["C12"].value == "L   L   W   L   D"  # last 5, most recent first
    assert sm["C38"].value == "Archie" and sm["C39"].value == "Archie, Max"
    assert sm["A16"].value == '=IF(Squad!$A2="","",Squad!$A2)'


def test_export_is_coaches_only(client: TestClient, db, demo: DemoSeason):
    make_user(db, "viewer", (UserRole.VIEWER, RoleScope.TEAM, demo.team_season.club_team_id))
    db.commit()
    login_as(client, "viewer")
    assert (
        client.get(f"{API}/team-seasons/{demo.team_season.id}/reports/season.xlsx").status_code
        == 403
    )


SOFFICE = shutil.which("soffice") or (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists()
    else None
)


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed")
def test_formulas_recalculate_to_the_apps_numbers(
    auth_client: TestClient, demo: DemoSeason, tmp_path
):
    """Recalculate in LibreOffice and compare the Summary with the stats service."""
    src = tmp_path / "season.xlsx"
    src.write_bytes(_download(auth_client, demo.team_season.id))
    out = tmp_path / "out"
    env = {**os.environ, "HOME": str(tmp_path)}  # fresh profile so soffice never prompts
    subprocess.run(
        [SOFFICE, "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
        check=True,
        capture_output=True,
        timeout=180,
        env=env,
    )
    wb = load_workbook(out / "season.xlsx", data_only=True)
    s = wb["Summary"]
    assert [s.cell(6, c).value for c in range(1, 8)] == [6, 2, 1, 3, 12, 12, 0]
    assert round(s["H6"].value, 3) == round(2 / 6, 3)
    assert [s.cell(10, c).value for c in range(1, 8)] == [4, 2, 1, 1, 11, 6, 5]
    squad = {
        s.cell(r, 1).value: [s.cell(r, c).value for c in range(2, 8)]
        for r in range(16, 36)
        if s.cell(r, 1).value
    }
    assert squad["Archie"] == [6, 5, 3, 0.83, 2, 0]
    assert squad["Max"] == [6, 2, 3, 0.33, 1, 2]
    assert squad["Teddy"] == [2, 0, 0, 0, 0, 0]
    assert squad["Kayson"] == [5, 1, 0, 0.2, 0, 1]
    assert s["D38"].value == "5 goals" and s["D39"].value == "3 assists"
    fx = wb["Fixtures"]
    assert [fx.cell(r, 7).value for r in range(2, 8)] == ["W", "D", "L", "W", "L", "L"]
    ap = wb["Appearances"]
    assert [ap.cell(r, 3).value for r in range(2, 8)] == [9, 7, 10, 9, 11, 8]
    for row in ap.iter_rows(min_row=2, max_row=8, min_col=2, max_col=2):
        assert row[0].value  # opposition looked up from Fixtures
    # Nothing evaluated to an error anywhere
    errors = ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.value not in errors, f"{ws.title}!{c.coordinate}={c.value}"
